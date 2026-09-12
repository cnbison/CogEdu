"""Phase 2 (2-E): 家长端/权限/报告导出真实进程灰度脚本 — 沿用 12.6/1-G 模式.

链路 (真实 uvicorn 进程, 登录全部走 HTTP — 2-0-3 无鉴权后门):
  开户 (直连灰度 DB: guardian×2 + student×1)
    → 家长 A 登录 → 发起绑定申请 (view_progress + download_report)
    → 学生登录 → 确认申请
    → 学生答题 ×3 (产生报告数据)
    → 家长 A roster/overview 可见
    → 家长 A 下载周报 → docx 重开 + 内容打印 (人工复核可读性/实用性)
    → 学生撤销授权 → 家长 A 下一请求立即 403
    → 家长 B (未关联) 访问 → 403

输出: 全链路结果打印到 stdout; 报告的段落/表格内容完整打印 —
人工检查版式/数字可读性 (2-E "不只是跑通不报错")。

用法:
  ECOS_DB_PATH=/tmp/canary_p2.db python scripts/canary_phase2_parent.py
不需要 LLM key (答题链路 correct 直接给定, 不调 judge)。
"""
from __future__ import annotations

import io
import json
import os
import subprocess
import sys
import tempfile
import time
import urllib.request
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

PORT = int(os.environ.get("CANARY_PORT", "5196"))
BASE = f"http://127.0.0.1:{PORT}"

CANARY_PASSWORD = "canary-password-2026"
SID = "canary_p2_stu"
GUARDIAN_A, GUARDIAN_B = "canary_p2_ga", "canary_p2_gb"

_STEP = 0


def req(method: str, path: str, body: dict | None = None, token: str | None = None):
    data = json.dumps(body).encode() if body is not None else None
    headers = {"Content-Type": "application/json"}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    r = urllib.request.Request(BASE + path, data=data, method=method, headers=headers)
    try:
        with urllib.request.urlopen(r, timeout=120) as resp:
            return resp.status, json.loads(resp.read())
    except urllib.error.HTTPError as e:
        body = e.read()
        try:
            return e.code, json.loads(body)
        except Exception:
            return e.code, {}


def download(method: str, path: str, token: str) -> tuple[int, bytes]:
    r = urllib.request.Request(
        BASE + path, method=method, headers={"Authorization": f"Bearer {token}"}
    )
    try:
        with urllib.request.urlopen(r, timeout=120) as resp:
            return resp.status, resp.read()
    except urllib.error.HTTPError as e:
        return e.code, b""


def ok(label: str, cond: bool, detail: str = "") -> None:
    global _STEP
    _STEP += 1
    mark = "✅" if cond else "❌"
    print(f"  {mark} [{_STEP:>2}] {label}" + (f"  ({detail})" if detail else ""))
    if not cond:
        raise SystemExit(f"[canary] 步骤失败: {label} {detail}")


def provision(db_path: str) -> None:
    """灰度开户 (直连 DB; v1 无注册端点, 与 1-G canary 同口径)."""
    os.environ["ECOS_DB_PATH"] = db_path
    from cogedu.persistence.db import get_db

    get_db().init_schema()
    from web.api import auth as auth_service

    accounts = [
        (GUARDIAN_A, "guardian", None),
        (GUARDIAN_B, "guardian", None),
        ("canary_p2_stu_acc", "student", SID),
    ]
    for username, role, sid in accounts:
        try:
            auth_service.create_user(
                username=username, password=CANARY_PASSWORD, role=sid and "student" or role,
                display_name=username, learning_student_id=sid,
            )
        except Exception as e:
            if "已存在" not in str(e):
                raise


def login(username: str) -> str:
    status, body = req("POST", "/api/auth/login", {
        "username": username, "password": CANARY_PASSWORD,
    })
    assert status == 200, f"登录失败 {username}: {status} {body}"
    return body["token"]


def print_report_docx(data: bytes) -> None:
    """docx 内容打印 (人工复核报告可读性 — 2-E 验收点)."""
    from docx import Document

    doc = Document(io.BytesIO(data))
    print("\n  ────────── 报告内容 (人工复核可读性) ──────────")
    for p in doc.paragraphs:
        if p.text.strip():
            print(f"  | {p.text}")
    for t in doc.tables:
        print("  | [表格]")
        for row in t.rows:
            print("  |   " + " | ".join(c.text for c in row.cells))
    print("  ────────────────────────────────────────────────")


def main() -> int:
    tmp_db = os.environ.get("ECOS_DB_PATH") or os.path.join(
        tempfile.mkdtemp(prefix="canary_p2_"), "canary.db"
    )
    provision(tmp_db)
    env = dict(os.environ, ECOS_DB_PATH=tmp_db)
    print(f"[canary] db={tmp_db}, port={PORT}")
    proc = subprocess.Popen(
        [sys.executable, "-m", "uvicorn", "web.api.app:app",
         "--host", "127.0.0.1", "--port", str(PORT)],
        env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
    )
    try:
        for _ in range(60):
            try:
                req("GET", "/api/version")
                print("[canary] server up")
                break
            except Exception:
                time.sleep(1)
        else:
            print("[canary] server failed to start")
            return 1

        ga = login(GUARDIAN_A)
        gb = login(GUARDIAN_B)
        stu = login("canary_p2_stu_acc")

        print("\n[链路] 申请 → 确认 → 数据可见")
        status, body = req("POST", "/api/guardian/links", {
            "learner_username": "canary_p2_stu_acc",
            "permissions": ["view_progress", "download_report"],
        }, token=ga)
        ok("家长 A 发起绑定申请", status == 200 and body["link"]["status"] == "pending")
        link_id = body["link"]["link_id"]

        status, body = req("GET", "/api/student/guardian-links", token=stu)
        ok("学生端可见待确认申请", status == 200
           and any(x["link_id"] == link_id for x in body["links"]))

        status, body = req("POST", f"/api/student/guardian-links/{link_id}/confirm", token=stu)
        ok("学生确认 → active", status == 200 and body["link"]["status"] == "active")

        print("\n[链路] 学生答题 → 家长看数据")
        for i in range(3):
            status, _ = req("POST", "/api/answer", {
                "student_id": SID, "problem_id": f"CANARY-P2-Q{i + 1}",
                "skill_id": "math.quadratic",
                "correct": i < 2, "score": 1.0 if i < 2 else 0.0,
                "bloom_layer": "L3",
            }, token=stu)
            ok(f"学生答题 Q{i + 1}", status == 200)

        status, body = req("GET", "/api/parent/students", token=ga)
        ok("家长 A roster 可见关联学生", status == 200
           and [s["student_id"] for s in body["students"]] == [SID])

        status, body = req("GET", f"/api/parent/students/{SID}/overview", token=ga)
        ok("家长 A overview 可见 (view_progress)", status == 200 and "five_d" in body)

        print("\n[链路] 报告下载 + 人工复核")
        status, data = download("GET", f"/api/parent/students/{SID}/report?period=week", ga)
        ok("家长 A 下载周报 (download_report)", status == 200 and data[:2] == b"PK")
        print_report_docx(data)

        print("\n[链路] 权限边界 (2-E 验收点)")
        status, _ = req("GET", f"/api/parent/students/{SID}/overview", token=gb)
        ok("家长 B (未关联) 访问 overview → 403", status == 403)

        status, _ = download("GET", f"/api/parent/students/{SID}/report?period=week", gb)
        ok("家长 B (未授权) 下载报告 → 403", status == 403)

        status, _ = req("DELETE", f"/api/student/guardian-links/{link_id}", token=stu)
        ok("学生撤销授权", status == 200)

        status, _ = req("GET", f"/api/parent/students/{SID}/overview", token=ga)
        ok("撤销后家长 A 下一请求立即 403", status == 403)

        status, _ = download("GET", f"/api/parent/students/{SID}/report?period=week", ga)
        ok("撤销后家长 A 下载报告立即 403", status == 403)

        print(f"\n[canary] 全部 {_STEP} 步通过 ✅")
        print("[canary] 待人工复核: 上方报告内容的版式/数字/文字可读性")
        return 0
    finally:
        proc.terminate()
        proc.wait(timeout=10)


if __name__ == "__main__":
    raise SystemExit(main())
